"""
Парсер расписания из xlsx-файла для группы ХТ-344.
Исправленная версия с корректным извлечением данных.
"""

import openpyxl
import json
import re
import sys
import os


def hours_to_pair(hours_str):
    """Преобразует академические часы в номер пары."""
    m = re.match(r'(\d+)\s*[-–]\s*(\d+)', str(hours_str))
    if m:
        first_hour = int(m.group(1))
        return str((first_hour + 1) // 2)
    return None


def find_group_column(ws, target_group="ХТ-344"):
    """Находит колонку и строку, где указана целевая группа."""
    for row in ws.iter_rows(min_row=1, max_row=50):
        for cell in row:
            if cell.value and target_group in str(cell.value).strip():
                return cell.column, cell.row
    return None, None


def find_next_group_column(ws, row_num, start_col):
    """Находит колонку следующей группы после целевой."""
    pattern = re.compile(r'(ХТ|РХТ)-\d+')
    for cell in ws[row_num]:
        if cell.column > start_col and cell.value:
            if pattern.search(str(cell.value)):
                return cell.column
    return None


def is_room(text):
    """Проверяет, является ли текст аудиторией."""
    if not text:
        return False
    text = str(text).strip()
    # Аудитории: Б-313, В 1302, 406, 345, 348, зал ГУК, Спортивный зал, Т 301
    patterns = [
        r'^[БВТ]\s?[-\s]?\d+',  # Б-313, В 1302, Т 301
        r'^\d{3,4}$',  # 406, 345, 348
        r'^зал',  # зал ГУК
        r'^спортивный',  # Спортивный зал
        r'^ауд',  # ауд.
    ]
    return any(re.match(p, text, re.IGNORECASE) for p in patterns)


def is_teacher(text):
    """Проверяет, является ли текст преподавателем."""
    if not text:
        return False
    text = str(text).strip()
    # ФИО: Фамилия И.О. или доц. Фамилия И.О. или проф. Фамилия И.О.
    patterns = [
        r'[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.?',  # Фамилия И.О.
        r'(доц\.|проф\.|ст\.преп\.)\s*[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.',  # доц. Фамилия И.О.
    ]
    return any(re.search(p, text) for p in patterns)


def parse_schedule(filename, target_group="ХТ-344"):
    if filename.lower().endswith('.xls') and not filename.lower().endswith('.xlsx'):
        print("❌ Ошибка: Файл имеет старый формат .xls!")
        print("Откройте файл в Excel/LibreOffice/Numbers и сделайте 'Сохранить как...' → .xlsx")
        sys.exit(1)

    print(f" Открываю файл: {filename}")
    wb = openpyxl.load_workbook(filename, data_only=True)

    if not wb.sheetnames:
        print("❌ В файле нет листов!")
        return None

    ws = wb.active if wb.active else wb[wb.sheetnames[0]]
    print(f"📄 Читаем лист: '{ws.title}'")

    # 1. Найти колонку группы
    group_col, group_row = find_group_column(ws, target_group)
    if not group_col:
        print(f"❌ Группа {target_group} не найдена в файле!")
        return None

    print(f"✅ Группа {target_group} найдена: колонка {group_col}, строка {group_row}")

    # 2. Определить диапазон колонок для группы
    next_col = find_next_group_column(ws, group_row, group_col)
    if next_col:
        col_start, col_end = group_col, next_col - 1
    else:
        col_start, col_end = group_col, group_col + 4

    print(f"📊 Диапазон колонок для группы: {col_start}–{col_end}")

    # 3. Парсинг
    DAYS_UPPER = ["ПОНЕДЕЛЬНИК", "ВТОРНИК", "СРЕДА", "ЧЕТВЕРГ", "ПЯТНИЦА", "СУББОТА"]
    DAYS_NORM = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота"]

    schedule = {
        "week1": {day: {} for day in DAYS_NORM},
        "week2": {day: {} for day in DAYS_NORM}
    }
    current_week = "week1"
    current_day = None
    current_pair = None

    # Собираем все строки в список для анализа
    rows_data = []
    for row in ws.iter_rows(min_row=group_row + 1, max_row=ws.max_row):
        values = {cell.column: cell.value for cell in row if cell.value is not None}
        rows_data.append(values)

    # Анализируем строки
    i = 0
    while i < len(rows_data):
        values = rows_data[i]

        # Проверяем переход ко второй неделе
        for col, val in values.items():
            val_str = str(val).strip()
            if target_group in val_str and col != group_col:
                if current_week == "week1":
                    current_week = "week2"
                    current_day = None
                    current_pair = None
                    print("🔄 Переход ко второй неделе")
                break

        # Ищем день недели
        for col, val in values.items():
            val_upper = str(val).strip().upper()
            if val_upper in DAYS_UPPER:
                idx = DAYS_UPPER.index(val_upper)
                current_day = DAYS_NORM[idx]
                current_pair = None
                break

        # Ищем академические часы
        for col, val in values.items():
            pair_num = hours_to_pair(val)
            if pair_num:
                current_pair = pair_num
                break

        # Если нашли день и пару, извлекаем данные
        if current_day and current_pair:
            # Собираем все данные из колонок группы для этой строки и следующих
            all_data = []
            j = i
            while j < len(rows_data) and j < i + 5:  # Берем до 5 строк
                row_vals = rows_data[j]
                for c in range(col_start, col_end + 1):
                    if c in row_vals:
                        text = str(row_vals[c]).strip()
                        if text and text != 'None':
                            all_data.append(text)
                j += 1

            if all_data:
                entry = {
                    "subject": "",
                    "teacher": "",
                    "room": ""
                }

                # Распределяем данные по полям
                for item in all_data:
                    if is_room(item):
                        if not entry["room"]:
                            entry["room"] = item
                    elif is_teacher(item):
                        if not entry["teacher"]:
                            entry["teacher"] = item
                    else:
                        # Это предмет (если еще не заполнен)
                        if not entry["subject"]:
                            entry["subject"] = item

                schedule[current_week][current_day][current_pair] = entry

        i += 1

    # 4. Очистка пустых дней
    for week in schedule:
        schedule[week] = {
            day: pairs for day, pairs in schedule[week].items() if pairs
        }

    result = {
        "group": target_group,
        "schedule": schedule
    }

    # 5. Сохранение
    out_path = "schedule.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n✅ Расписание сохранено в {out_path}")
    print(f"⚠️  Обязательно проверьте {out_path} вручную!")

    # Вывод статистики
    for week in schedule:
        print(f"\n📆 {week}:")
        for day, pairs in schedule[week].items():
            if pairs:
                print(f"  {day}:")
                for pair_num, data in sorted(pairs.items(), key=lambda x: int(x[0])):
                    print(f"    {pair_num} пара: {data.get('subject', '—')} | {data.get('teacher', '—')} | {data.get('room', '—')}")

    return result


if __name__ == "__main__":
    if len(sys.argv) > 1:
        fname = sys.argv[1]
    else:
        fname = "ОН_ХТФ_3 курс.xlsx"

    if not os.path.exists(fname):
        print(f"❌ Файл '{fname}' не найден!")
        sys.exit(1)

    parse_schedule(fname, "ХТ-344")